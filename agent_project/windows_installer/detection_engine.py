#!/usr/bin/env python3
"""
SecretSnipe Detection Engine - Full V1 Scanner Capabilities
Cross-platform detection engine for Windows/Linux agents

This module provides:
- Signature-based pattern detection (from signatures.json)
- OCR support for images (optional - Tesseract/EasyOCR)
- PDF text extraction (PyMuPDF/fitz)
- Excel extraction (openpyxl, xlrd)
- Word document extraction (python-docx)
- ZIP archive extraction
- Credit card validation (Luhn algorithm)
- SSN/AWS/JWT validation
- Entropy detection for tokens
- False positive reduction

Version: 1.0.0
"""

import os
import re
import json
import math
import logging
import hashlib
import tempfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Any
from datetime import datetime

# ============================================================================
# OPTIONAL IMPORTS - Features enabled based on availability
# ============================================================================

# PDF Support
try:
    import fitz  # PyMuPDF
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

# Excel Support
try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    import xlrd
    XLRD_AVAILABLE = True
except ImportError:
    XLRD_AVAILABLE = False

# Word Document Support
try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# OCR Support - Tesseract
try:
    import pytesseract
    from PIL import Image
    TESSERACT_AVAILABLE = True
except ImportError:
    TESSERACT_AVAILABLE = False

# OCR Support - EasyOCR (fallback, more accurate but higher memory)
try:
    import easyocr
    EASYOCR_AVAILABLE = True
except ImportError:
    EASYOCR_AVAILABLE = False

# ZIP Support
from zipfile import ZipFile, BadZipFile

# ============================================================================
# DETECTION ENGINE
# ============================================================================

class DetectionEngine:
    """
    Full-featured secret detection engine with V1 scanner parity.
    
    Features:
    - Signature-based detection with compiled regex
    - Multi-format file support (text, PDF, Excel, Word, images, archives)
    - Validation functions (Luhn, SSN, AWS, JWT, entropy)
    - False positive reduction
    - Configurable OCR support
    """
    
    # Default file extensions for plain text
    TEXT_EXTENSIONS = {
        '.txt', '.py', '.js', '.ts', '.java', '.cpp', '.c', '.h',
        '.php', '.rb', '.go', '.rs', '.swift', '.kt', '.scala',
        '.clj', '.hs', '.ml', '.md', '.json', '.xml', '.yaml',
        '.yml', '.toml', '.ini', '.cfg', '.conf', '.properties',
        '.sh', '.bat', '.ps1', '.sql', '.html', '.css', '.scss',
        '.env', '.config', '.htaccess', '.gitignore', '.dockerignore',
        '.cs', '.vb', '.fs', '.dart', '.lua', '.perl', '.pl',
        '.r', '.R', '.m', '.mm', '.gradle', '.cmake', '.makefile',
        '.tf', '.hcl', '.pem', '.key', '.crt', '.cer'
    }
    
    # Image extensions for OCR
    IMAGE_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.tif', '.gif'}
    
    # Default signatures (embedded, can be overridden with signatures.json)
    DEFAULT_SIGNATURES = [
        {
            "name": "Hardcoded Password",
            "regex": r"(?i)(?:password|passwd|pwd|secret|token)\s*[:=]\s*(?!true|false|null|none|undefined|empty|changeme|placeholder|example|\${|\[|\{\{|\b\w+\b\.|\s*['\"]?$)['\"]?([^'\"\\s,;$%<>\\n\\r]{8,64})['\"]?",
            "severity": "HIGH",
            "validate_entropy": True
        },
        {
            "name": "API Key",
            "regex": r"(?i)api[_-]?key\s*[:=]\s*['\"]?(?!your|example|placeholder|changeme|xxx)([a-zA-Z0-9-_.]{20,64})['\"]?",
            "severity": "HIGH",
            "validate_entropy": True
        },
        {
            "name": "Private Key",
            "regex": r"-----BEGIN (RSA|EC|OPENSSH|PGP|DSA|ENCRYPTED) PRIVATE KEY-----[\s\S]*?-----END \1 PRIVATE KEY-----",
            "severity": "CRITICAL"
        },
        {
            "name": "Database Connection String",
            "regex": r"(?i)(?:jdbc|mysql|postgresql|mongodb(?:\+srv)?|redis|mssql|oracle)://(?:[\w.-]+:[^@\s]+@)[\w.:/-]+",
            "severity": "CRITICAL"
        },
        {
            "name": "AWS Access Key ID",
            "regex": r"(?:^|[^A-Z0-9])AKIA[0-9A-Z]{16}(?:[^A-Z0-9]|$)",
            "severity": "HIGH",
            "validate_aws": True
        },
        {
            "name": "AWS Secret Access Key",
            "regex": r"(?i)(?:aws[_-]?secret[_-]?(?:access[_-]?)?key|aws_secret)\s*[:=]\s*['\"]?([A-Za-z0-9/+=]{40})['\"]?",
            "severity": "CRITICAL"
        },
        {
            "name": "Generic Auth Token",
            "regex": r"(?i)bearer\s+[a-zA-Z0-9-._~+/]{20,}=*",
            "severity": "HIGH",
            "validate_entropy": True
        },
        {
            "name": "Credit Card Number",
            "regex": r"(?i)(?:card|credit|payment|visa|mastercard|amex|discover|pan|ccn|cardnum)[^\n]{0,40}\b(?:4[0-9]{3}[- ]?[0-9]{4}[- ]?[0-9]{4}[- ]?[0-9]{4}|5[1-5][0-9]{2}[- ]?[0-9]{4}[- ]?[0-9]{4}[- ]?[0-9]{4}|3[47][0-9]{2}[- ]?[0-9]{6}[- ]?[0-9]{5}|6(?:011|5[0-9]{2})[- ]?[0-9]{4}[- ]?[0-9]{4}[- ]?[0-9]{4})\b",
            "severity": "CRITICAL",
            "validate_credit_card": True
        },
        {
            "name": "Social Security Number",
            "regex": r"(?i)(?:ssn|social[\s_-]*security|soc[\s_-]*sec)\s*(?:number|num|no|#)?\s*[:=]?\s*['\"]?(\d{3}-\d{2}-\d{4})['\"]?",
            "severity": "HIGH",
            "validate_ssn": True
        },
        {
            "name": "JWT Token",
            "regex": r"eyJ[A-Za-z0-9-_]{20,}\.eyJ[A-Za-z0-9-_]{20,}\.[A-Za-z0-9-_.+/]{20,}",
            "severity": "HIGH",
            "validate_jwt": True
        },
        {
            "name": "GitHub Token",
            "regex": r"(?:ghp|gho|ghu|ghs|ghr)_[0-9A-Za-z]{36}",
            "severity": "CRITICAL"
        },
        {
            "name": "GitLab Token",
            "regex": r"glpat-[0-9A-Za-z_-]{20}",
            "severity": "CRITICAL"
        },
        {
            "name": "Slack Token",
            "regex": r"xox[baprs]-[0-9A-Za-z-]{10,}",
            "severity": "HIGH",
            "validate_entropy": True
        },
        {
            "name": "Slack Webhook",
            "regex": r"https://hooks\.slack\.com/services/T[a-zA-Z0-9_]{8,}/B[a-zA-Z0-9_]{8,}/[a-zA-Z0-9_]{20,}",
            "severity": "HIGH"
        },
        {
            "name": "Azure Storage Key",
            "regex": r"(?i)(?:account[_-]?key|storage[_-]?key)\s*[:=]\s*['\"]?([A-Za-z0-9+/]{86}==)['\"]?",
            "severity": "CRITICAL"
        },
        {
            "name": "Azure Connection String",
            "regex": r"(?i)DefaultEndpointsProtocol=https?;AccountName=[^;]+;AccountKey=[A-Za-z0-9+/]{86}==",
            "severity": "CRITICAL"
        },
        {
            "name": "Google API Key",
            "regex": r"AIza[0-9A-Za-z_-]{35}",
            "severity": "HIGH"
        },
        {
            "name": "Stripe API Key",
            "regex": r"sk_live_[0-9a-zA-Z]{24,}",
            "severity": "CRITICAL"
        },
        {
            "name": "SendGrid API Key",
            "regex": r"SG\.[a-zA-Z0-9_-]{22}\.[a-zA-Z0-9_-]{43}",
            "severity": "HIGH",
            "validate_entropy": True
        },
        {
            "name": "Twilio API Key",
            "regex": r"SK[0-9a-fA-F]{32}",
            "severity": "HIGH"
        },
        {
            "name": "NPM Token",
            "regex": r"npm_[A-Za-z0-9]{36}",
            "severity": "HIGH"
        },
        {
            "name": "RSA Private Key",
            "regex": r"-----BEGIN RSA PRIVATE KEY-----[\s\S]{100,}-----END RSA PRIVATE KEY-----",
            "severity": "CRITICAL"
        },
        {
            "name": "SSH Private Key",
            "regex": r"-----BEGIN OPENSSH PRIVATE KEY-----[\s\S]{100,}-----END OPENSSH PRIVATE KEY-----",
            "severity": "CRITICAL"
        },
        {
            "name": "PGP Private Key",
            "regex": r"-----BEGIN PGP PRIVATE KEY BLOCK-----[\s\S]{100,}-----END PGP PRIVATE KEY BLOCK-----",
            "severity": "CRITICAL"
        },
        {
            "name": "High Entropy String",
            "regex": r"['\"][a-zA-Z0-9+/=_-]{32,}['\"]",
            "severity": "MEDIUM",
            "validate_entropy": True,
            "min_entropy": 4.0,
            "low_risk": True
        }
    ]
    
    # Known test credit card numbers
    TEST_CREDIT_CARDS = {
        '4242424242424242', '4000056655665556', '5555555555554444',
        '2223003122003222', '5200828282828210', '378282246310005',
        '371449635398431', '6011111111111117', '6011000990139424',
        '4111111111111111', '4012888888881881', '4009348888881881',
        '4012000033330026', '4012000077777777', '370000000000002',
        '6011000000000012', '0000000000000000', '1111111111111111',
        '1234567890123456'
    }
    
    def __init__(self, logger: logging.Logger = None, config: Dict = None):
        """Initialize detection engine"""
        self.logger = logger or logging.getLogger("DetectionEngine")
        self.config = config or {}
        
        # Signature list (loaded or default)
        self.signatures = []
        self._compiled_patterns = []
        
        # OCR reader (lazy loaded)
        self._ocr_reader = None
        self._ocr_count = 0
        
        # Configuration
        self.enable_ocr = self.config.get("enable_ocr", True)
        self.enable_pdf = self.config.get("enable_pdf", True)
        self.enable_office = self.config.get("enable_office", True)
        self.max_file_size_mb = self.config.get("max_file_size_mb", 50)
        self.max_ocr_file_size_mb = self.config.get("max_ocr_file_size_mb", 5.0)
        self.ocr_engine = self.config.get("ocr_engine", "tesseract")  # tesseract or easyocr
        
        # Load signatures
        self._load_signatures()
        
        # Log capabilities
        self._log_capabilities()
    
    def _log_capabilities(self):
        """Log available detection capabilities"""
        caps = []
        caps.append("signatures")
        if PDF_AVAILABLE and self.enable_pdf:
            caps.append("PDF")
        if (OPENPYXL_AVAILABLE or XLRD_AVAILABLE) and self.enable_office:
            caps.append("Excel")
        if DOCX_AVAILABLE and self.enable_office:
            caps.append("Word")
        if self.enable_ocr:
            if TESSERACT_AVAILABLE:
                caps.append("OCR-Tesseract")
            elif EASYOCR_AVAILABLE:
                caps.append("OCR-EasyOCR")
        self.logger.info(f"🔍 Detection Engine initialized: {', '.join(caps)}")
    
    def _load_signatures(self):
        """Load signatures from file or use defaults"""
        signatures_file = self.config.get("signatures_file", "signatures.json")
        
        # Try to load from file
        if os.path.exists(signatures_file):
            try:
                with open(signatures_file, 'r') as f:
                    raw_sigs = json.load(f)
                self.signatures = raw_sigs
                self.logger.info(f"📝 Loaded {len(raw_sigs)} signatures from {signatures_file}")
            except Exception as e:
                self.logger.warning(f"Failed to load {signatures_file}: {e}, using defaults")
                self.signatures = self.DEFAULT_SIGNATURES
        else:
            self.signatures = self.DEFAULT_SIGNATURES
            self.logger.info(f"📝 Using {len(self.signatures)} default signatures")
        
        # Compile patterns
        self._compile_patterns()
    
    def _compile_patterns(self):
        """Pre-compile regex patterns for performance"""
        self._compiled_patterns = []
        for sig in self.signatures:
            try:
                pattern_str = sig.get("regex") or sig.get("pattern", "")
                if not pattern_str:
                    continue
                
                compiled = re.compile(pattern_str, re.IGNORECASE | re.MULTILINE)
                self._compiled_patterns.append({
                    "name": sig.get("name", "Unknown"),
                    "pattern": compiled,
                    "severity": sig.get("severity", "MEDIUM").upper(),
                    "description": sig.get("description", ""),
                    "validate_credit_card": sig.get("validate_credit_card", False),
                    "validate_ssn": sig.get("validate_ssn", False),
                    "validate_aws": sig.get("validate_aws", False),
                    "validate_jwt": sig.get("validate_jwt", False),
                    "validate_entropy": sig.get("validate_entropy", False),
                    "min_entropy": sig.get("min_entropy", 3.5),
                    "low_risk": sig.get("low_risk", False)
                })
            except re.error as e:
                self.logger.warning(f"Invalid regex for {sig.get('name', 'unknown')}: {e}")
    
    # ========================================================================
    # TEXT EXTRACTION
    # ========================================================================
    
    def extract_text(self, file_path: Path) -> str:
        """Extract text from various file formats"""
        file_ext = file_path.suffix.lower()
        
        try:
            # Check file size
            file_size_mb = file_path.stat().st_size / (1024 * 1024)
            if file_size_mb > self.max_file_size_mb:
                self.logger.debug(f"File too large ({file_size_mb:.1f}MB): {file_path}")
                return ""
            
            # Plain text files
            if file_ext in self.TEXT_EXTENSIONS:
                return self._extract_text_file(file_path)
            
            # PDF files
            elif file_ext == '.pdf' and PDF_AVAILABLE and self.enable_pdf:
                return self._extract_pdf(file_path)
            
            # Excel files
            elif file_ext in ['.xlsx', '.xls'] and self.enable_office:
                return self._extract_excel(file_path)
            
            # Word documents
            elif file_ext in ['.docx'] and DOCX_AVAILABLE and self.enable_office:
                return self._extract_word(file_path)
            
            # Images with OCR
            elif file_ext in self.IMAGE_EXTENSIONS and self.enable_ocr:
                return self._extract_image_ocr(file_path)
            
            # ZIP archives
            elif file_ext == '.zip':
                return self._extract_zip(file_path)
            
            else:
                self.logger.debug(f"Unsupported file type: {file_ext}")
                return ""
                
        except Exception as e:
            self.logger.warning(f"Error extracting text from {file_path}: {e}")
            return ""
    
    def _extract_text_file(self, file_path: Path) -> str:
        """Extract text from plain text files with encoding detection"""
        try:
            raw_bytes = file_path.read_bytes()
            
            # Check for BOM markers
            if raw_bytes.startswith(b'\xff\xfe'):  # UTF-16 LE
                return raw_bytes.decode('utf-16-le', errors='ignore')
            elif raw_bytes.startswith(b'\xfe\xff'):  # UTF-16 BE
                return raw_bytes.decode('utf-16-be', errors='ignore')
            elif raw_bytes.startswith(b'\xef\xbb\xbf'):  # UTF-8 BOM
                return raw_bytes[3:].decode('utf-8', errors='ignore')
            else:
                # Try UTF-8 first
                try:
                    return raw_bytes.decode('utf-8')
                except UnicodeDecodeError:
                    # Check for UTF-16 without BOM (null bytes)
                    if b'\x00' in raw_bytes[:100]:
                        return raw_bytes.decode('utf-16-le', errors='ignore')
                    else:
                        return raw_bytes.decode('latin-1')
        except Exception as e:
            self.logger.warning(f"Text extraction failed for {file_path}: {e}")
            return ""
    
    def _extract_pdf(self, file_path: Path) -> str:
        """Extract text from PDF files"""
        if not PDF_AVAILABLE:
            return ""
        
        try:
            text = ""
            with fitz.open(str(file_path)) as doc:
                for page in doc:
                    text += page.get_text()
            return text
        except Exception as e:
            self.logger.warning(f"PDF extraction failed for {file_path}: {e}")
            return ""
    
    def _extract_excel(self, file_path: Path) -> str:
        """Extract text from Excel files"""
        text = ""
        file_ext = file_path.suffix.lower()
        
        try:
            if file_ext == '.xlsx' and OPENPYXL_AVAILABLE:
                workbook = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
                for sheet in workbook.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        text += ' '.join(str(cell) for cell in row if cell) + '\n'
                workbook.close()
            
            elif file_ext == '.xls' and XLRD_AVAILABLE:
                workbook = xlrd.open_workbook(str(file_path))
                for sheet in workbook.sheets():
                    for row_idx in range(sheet.nrows):
                        row = sheet.row_values(row_idx)
                        text += ' '.join(str(cell) for cell in row if cell) + '\n'
                        
        except Exception as e:
            self.logger.warning(f"Excel extraction failed for {file_path}: {e}")
        
        return text
    
    def _extract_word(self, file_path: Path) -> str:
        """Extract text from Word documents"""
        if not DOCX_AVAILABLE:
            return ""
        
        try:
            doc = Document(str(file_path))
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + '\n'
            return text
        except Exception as e:
            self.logger.warning(f"Word extraction failed for {file_path}: {e}")
            return ""
    
    def _extract_image_ocr(self, file_path: Path) -> str:
        """Extract text from images using OCR"""
        # Check file size
        try:
            file_size_mb = file_path.stat().st_size / (1024 * 1024)
            if file_size_mb > self.max_ocr_file_size_mb:
                self.logger.debug(f"Image too large for OCR ({file_size_mb:.1f}MB): {file_path}")
                return ""
        except:
            return ""
        
        # Try Tesseract first (lower memory)
        if TESSERACT_AVAILABLE and self.ocr_engine == "tesseract":
            try:
                img = Image.open(str(file_path))
                text = pytesseract.image_to_string(img, timeout=30)
                img.close()
                self._ocr_count += 1
                return text.strip()
            except Exception as e:
                self.logger.warning(f"Tesseract OCR failed for {file_path}: {e}")
        
        # Fallback to EasyOCR
        if EASYOCR_AVAILABLE:
            try:
                if self._ocr_reader is None:
                    self.logger.info("Initializing EasyOCR reader...")
                    self._ocr_reader = easyocr.Reader(['en'])
                
                results = self._ocr_reader.readtext(str(file_path))
                text = ' '.join([r[1] for r in results])
                self._ocr_count += 1
                return text.strip()
            except Exception as e:
                self.logger.warning(f"EasyOCR failed for {file_path}: {e}")
        
        return ""
    
    def _extract_zip(self, file_path: Path) -> str:
        """Extract text from ZIP archives"""
        text = ""
        try:
            with ZipFile(str(file_path), 'r') as zf:
                for info in zf.filelist:
                    if not info.is_dir() and info.file_size < 1024 * 1024:  # 1MB limit per file
                        try:
                            content = zf.read(info.filename).decode('utf-8', errors='ignore')
                            text += content + '\n'
                        except:
                            pass
        except BadZipFile:
            self.logger.debug(f"Invalid ZIP file: {file_path}")
        except Exception as e:
            self.logger.warning(f"ZIP extraction failed for {file_path}: {e}")
        
        return text
    
    # ========================================================================
    # VALIDATION FUNCTIONS
    # ========================================================================
    
    def luhn_checksum(self, card_number: str) -> bool:
        """Validate credit card using Luhn algorithm"""
        digits = ''.join(c for c in card_number if c.isdigit())
        
        if not digits or len(digits) < 13 or len(digits) > 19:
            return False
        
        digits_list = [int(d) for d in digits][::-1]
        checksum = 0
        for i, digit in enumerate(digits_list):
            if i % 2 == 1:
                digit *= 2
                if digit > 9:
                    digit -= 9
            checksum += digit
        
        return checksum % 10 == 0
    
    def is_test_credit_card(self, card_number: str) -> bool:
        """Check if card is a known test number"""
        digits = ''.join(c for c in card_number if c.isdigit())
        
        if len(digits) < 13:
            return False
        
        # Check known test cards
        if digits in self.TEST_CREDIT_CARDS:
            return True
        
        # Check repeating patterns
        if len(digits) == 16:
            groups = [digits[i:i+4] for i in range(0, 16, 4)]
            if len(set(groups)) == 1:
                return True
        
        # All same digits
        if len(set(digits)) == 1:
            return True
        
        return False
    
    def validate_ssn(self, ssn: str) -> Tuple[bool, str]:
        """Validate Social Security Number"""
        digits = ''.join(c for c in ssn if c.isdigit())
        
        if len(digits) != 9:
            return False, "Invalid length"
        
        area = int(digits[:3])
        group = int(digits[3:5])
        serial = int(digits[5:])
        
        if area == 0 or area == 666 or area >= 900:
            return False, "Invalid area number"
        if group == 0:
            return False, "Invalid group number"
        if serial == 0:
            return False, "Invalid serial number"
        
        # Known invalid SSNs
        known_invalid = {'078051120', '219099999', '457555462', '123456789',
                         '111111111', '222222222', '333333333', '444444444',
                         '555555555', '666666666', '777777777', '888888888', '999999999'}
        if digits in known_invalid:
            return False, "Known invalid SSN"
        
        return True, ""
    
    def validate_aws_key(self, key: str) -> Tuple[bool, str]:
        """Validate AWS Access Key ID"""
        key = key.strip()
        
        if len(key) != 20:
            return False, f"Invalid length: {len(key)}"
        
        if not key.startswith(('AKIA', 'ASIA', 'AIDA', 'AROA', 'AIPA', 'ANPA', 'ANVA', 'AGPA')):
            return False, "Invalid prefix"
        
        test_keys = {'AKIAIOSFODNN7EXAMPLE', 'AKIAI44QH8DHBEXAMPLE'}
        if key in test_keys:
            return False, "Known test key"
        
        return True, ""
    
    def validate_jwt(self, token: str) -> Tuple[bool, str]:
        """Validate JWT token structure"""
        import base64
        
        parts = token.split('.')
        if len(parts) != 3:
            return False, "Invalid JWT structure"
        
        try:
            # Decode header
            header_b64 = parts[0] + '=' * (4 - len(parts[0]) % 4)
            header = json.loads(base64.urlsafe_b64decode(header_b64))
            
            # Decode payload
            payload_b64 = parts[1] + '=' * (4 - len(parts[1]) % 4)
            payload = json.loads(base64.urlsafe_b64decode(payload_b64))
            
            # Check for test tokens
            sub = payload.get('sub', '')
            if sub in ['1234567890', 'test', 'example', 'user', 'admin']:
                return False, f"Likely test token (sub: {sub})"
            
            return True, ""
            
        except Exception as e:
            return False, f"Invalid JWT: {e}"
    
    def calculate_entropy(self, text: str) -> float:
        """Calculate Shannon entropy of text"""
        if len(text) < 10:
            return 0.0
        
        freq = {}
        for char in text:
            freq[char] = freq.get(char, 0) + 1
        
        entropy = 0.0
        for count in freq.values():
            p = count / len(text)
            if p > 0:
                entropy -= p * math.log2(p)
        
        return entropy
    
    def is_in_comment(self, context: str) -> bool:
        """Check if match is within a comment"""
        lines = context.split('\n')
        for line in lines:
            line = line.strip()
            if line.startswith('//') or line.startswith('#') or line.startswith('/*') or '*/' in line:
                return True
        return False
    
    # ========================================================================
    # SCANNING
    # ========================================================================
    
    def scan_text(self, text: str, file_path: str) -> List[Dict]:
        """Scan text for secrets using compiled signatures"""
        findings = []
        
        for sig in self._compiled_patterns:
            try:
                for match in sig["pattern"].finditer(text):
                    value = match.group(0)
                    
                    # Get context
                    start = max(0, match.start() - 50)
                    end = min(len(text), match.end() + 50)
                    context = text[start:end]
                    
                    # Validation
                    is_valid = True
                    validation_reason = ""
                    
                    # Check if in comment
                    if self.is_in_comment(context):
                        is_valid = False
                        validation_reason = "In comment"
                    
                    # Credit card validation
                    if is_valid and sig.get("validate_credit_card"):
                        if self.is_test_credit_card(value):
                            is_valid = False
                            validation_reason = "Test credit card"
                        elif not self.luhn_checksum(value):
                            is_valid = False
                            validation_reason = "Failed Luhn checksum"
                    
                    # SSN validation
                    if is_valid and sig.get("validate_ssn"):
                        ssn_valid, reason = self.validate_ssn(value)
                        if not ssn_valid:
                            is_valid = False
                            validation_reason = f"Invalid SSN: {reason}"
                    
                    # AWS key validation
                    if is_valid and sig.get("validate_aws"):
                        aws_valid, reason = self.validate_aws_key(value)
                        if not aws_valid:
                            is_valid = False
                            validation_reason = f"Invalid AWS key: {reason}"
                    
                    # JWT validation
                    if is_valid and sig.get("validate_jwt"):
                        jwt_valid, reason = self.validate_jwt(value)
                        if not jwt_valid:
                            is_valid = False
                            validation_reason = f"Invalid JWT: {reason}"
                    
                    # Entropy validation
                    if is_valid and sig.get("validate_entropy"):
                        min_entropy = sig.get("min_entropy", 3.5)
                        # Extract the actual secret value (group 1 if exists)
                        secret_value = match.group(1) if match.groups() else value
                        entropy = self.calculate_entropy(secret_value)
                        if entropy < min_entropy:
                            is_valid = False
                            validation_reason = f"Low entropy ({entropy:.2f} < {min_entropy})"
                    
                    # Mask the secret
                    matched = match.group(1) if match.groups() else value
                    if len(matched) > 8:
                        masked = matched[:4] + '*' * (len(matched) - 8) + matched[-4:]
                    else:
                        masked = '*' * len(matched)
                    
                    line_num = text[:match.start()].count('\n') + 1
                    
                    finding = {
                        "file": file_path,
                        "line": line_num,
                        "rule": sig["name"],
                        "severity": sig["severity"],
                        "match": masked,
                        "line_content": context.split('\n')[0][:200] if context else "",
                        "scanner": "detection_engine",
                        "is_valid": is_valid,
                        "validation_reason": validation_reason,
                        "confidence": 0.9 if is_valid else 0.5,
                        "timestamp": datetime.utcnow().isoformat()
                    }
                    findings.append(finding)
                    
                    if is_valid:
                        self.logger.debug(f"✅ Found: {sig['name']} at line {line_num} in {file_path}")
                    else:
                        self.logger.debug(f"⚠️ Filtered: {sig['name']} - {validation_reason}")
                        
            except Exception as e:
                self.logger.warning(f"Error scanning with {sig['name']}: {e}")
        
        return findings
    
    def scan_file(self, file_path: Path) -> List[Dict]:
        """Scan a single file for secrets"""
        # Extract text from file
        text = self.extract_text(file_path)
        
        if not text or len(text) < 10:
            return []
        
        # Scan text for secrets
        return self.scan_text(text, str(file_path))
    
    def get_capabilities(self) -> Dict[str, bool]:
        """Return current detection capabilities"""
        return {
            "signatures": len(self._compiled_patterns),
            "pdf": PDF_AVAILABLE and self.enable_pdf,
            "excel": (OPENPYXL_AVAILABLE or XLRD_AVAILABLE) and self.enable_office,
            "word": DOCX_AVAILABLE and self.enable_office,
            "ocr_tesseract": TESSERACT_AVAILABLE and self.enable_ocr,
            "ocr_easyocr": EASYOCR_AVAILABLE and self.enable_ocr,
            "zip": True
        }


# ============================================================================
# STANDALONE USAGE
# ============================================================================

if __name__ == "__main__":
    import argparse
    
    parser = argparse.ArgumentParser(description="SecretSnipe Detection Engine")
    parser.add_argument("path", help="File or directory to scan")
    parser.add_argument("--signatures", help="Path to signatures.json")
    parser.add_argument("--verbose", "-v", action="store_true")
    args = parser.parse_args()
    
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format='%(asctime)s - %(levelname)s - %(message)s'
    )
    
    config = {}
    if args.signatures:
        config["signatures_file"] = args.signatures
    
    engine = DetectionEngine(config=config)
    
    target = Path(args.path)
    if target.is_file():
        findings = engine.scan_file(target)
    else:
        findings = []
        for root, dirs, files in os.walk(target):
            for f in files:
                file_path = Path(root) / f
                findings.extend(engine.scan_file(file_path))
    
    print(f"\n{'='*60}")
    print(f"Found {len(findings)} potential secrets")
    print(f"{'='*60}\n")
    
    for finding in findings:
        if finding.get("is_valid", True):
            print(f"[{finding['severity']}] {finding['rule']}")
            print(f"  File: {finding['file']}:{finding['line']}")
            print(f"  Match: {finding['match']}")
            print()
